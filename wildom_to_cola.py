"""
Wildom Cola-rendeles -> PizzaMe_megrendelo sablon automatikus feltoltese.

Hasznalat:
    python3 wildom_to_cola.py <wildom_export.xlsx> <pizzame_sablon.xlsx> [uj_lap_nev]

- <wildom_export.xlsx>: a Wildom "Beszallitoi megrendelesek" -> "Excel 2007+"
  exporttal letoltott fajl (pl. "megrendeles-Coca Cola-YYYYMMDD-HHMM.xlsx").
- <pizzame_sablon.xlsx>: a PizzaMe_megrendelo fajl (a legutobbi allapot,
  amibol uj lapot masolunk).
- [uj_lap_nev]: opcionalis, ha nem adod meg, a mai datum lesz (YYYYMMDD).

A script megkeresi a legfrissebb lapot a sablonban (a legnagyobb, csak
szambol allo lapnevet), lemasolja azt uj lapkent, es a Wildom exportbol
kitolti a mennyisegeket a rogzitett termek- es uzlet-parositas alapjan.

A termek- es uzletlista karbantartasa: lasd a PRODUCT_MAP es COL_MAP
szotarakat lent, illetve a Wildom_Cola_parositas.xlsx fajlt a
dokumentaciohoz.
"""
import sys
import re
import copy
import datetime
import openpyxl

# --- Wildom termeknev -> PizzaMe sablon cikkszam ---
PRODUCT_MAP = {
    "Coca-Cola": 1188084,
    "Coca-Cola Zero": 609280,
    "Naturaqua Savas 0,5l": 1148406,
    "Naturaqua Mentes 0,5l": 295006,
    "Fuzetea Barack 0,5l": 1888013,
    "Fuzetea Citrom 0,5l": 1676459,
    "Cappy Narancs": 1523602,
    "Cappy Alma": 983327,
    "Kinley Gyömbér 0,5 L": 686007,
    "Coca-Cola 1L": 2456204,
    "Coca-Cola Zero 1L": 401274,
    "Coca-Cola 0,5L": 256553,
    "Coca-Cola Zero 0,5L ": 1307304,
    "Coca-Cola Koffeinmentes Zero 0,33l": 1180102,
    "Jack&Coke": 2462660,
    "Naturaqua Mentes 1,5l": 295217,
    "Naturaqua Savas 1,5l": 295118,
    "Cappy Ice Fruit Multivitamin 0,5L": 279045,
    "Kinley Pink 0,5l": 2474103,
}

# --- Sablon oszlop index (0-based, a fejlista sorrendje szerint) -> Wildom oszlop nev ---
COL_MAP_BY_NAME = {
    3: "Király", 4: "Erzsébet krt. 51(2)", 5: "Fashion", 6: "Bazilika",
    7: "Széll Kálmán Tér", 8: "Corvin", 9: "Károly krt.", 10: "Dob",
    11: "Blaha", 12: "Oktogon", 13: "Wesselényi", 15: "Astoria",
    16: "Erzsébet krt. 14. (13)", 17: "Gomba", 19: "Keleti",
    20: "Fővám Tér", 21: "Ferenciek tere", 22: "Jászai Mari Tér",
    23: "Etele Plaza", 24: "Óbuda", 25: "Kispest", 26: "Bosnyák tér",
    27: "Árkád", 28: "Kazinczy", 29: "Eleven Center", 30: "Törökvész",
    31: "GoBuda", 34: "D1 Buda", 35: "D2 Baross", 36: "Siófok",
    37: "Szeged", 39: "Pécs", 40: "Miskolc",
}
# Uj oszlopok (nincsenek meg a regi sablonban) - hozzuk lettre, ha hianyoznak
NEW_COLUMNS = {
    "Pizza Me Újpest": "Újpest",
    "Pizza Me K1 Westend": "K1 Westend",
}
# Ezeket a Wildom oszlopokat tudottan KIHAGYJUK (felhasznaloi donces alapjan)
EXCLUDED_WILDOM_COLS = {"Truck 1 Kamion", "Budapest Park", "PM Plázs", "Closed Móricz"}


def find_latest_sheet(wb):
    candidates = [s for s in wb.sheetnames if re.fullmatch(r"\d{6,8}", s)]
    if not candidates:
        return wb.sheetnames[0]
    return sorted(candidates, key=lambda s: int(s))[-1]


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    wildom_path, sablon_path = sys.argv[1], sys.argv[2]
    new_sheet_name = sys.argv[3] if len(sys.argv) > 3 else datetime.date.today().strftime("%Y%m%d")

    wb_w = openpyxl.load_workbook(wildom_path, data_only=True)
    ws_w = wb_w[wb_w.sheetnames[0]]
    w_headers = [c.value for c in ws_w[1]]
    w_idx = {name: i for i, name in enumerate(w_headers)}
    w_data = {row[0]: row for row in ws_w.iter_rows(min_row=2, max_row=ws_w.max_row, values_only=True) if row[0]}

    wb_p = openpyxl.load_workbook(sablon_path)
    src_name = find_latest_sheet(wb_p)
    src = wb_p[src_name]

    if new_sheet_name in wb_p.sheetnames:
        del wb_p[new_sheet_name]
    ws = wb_p.copy_worksheet(src)
    ws.title = new_sheet_name

    max_col = ws.max_column
    sample_header = ws.cell(row=5, column=4)
    sample_qty = ws.cell(row=6, column=4)

    # meglevo oszlopok fejlistaja a masolt lapon
    header_row = [ws.cell(row=5, column=c).value for c in range(1, max_col + 1)]

    col_map = {}
    for sablon_idx, wildom_name in COL_MAP_BY_NAME.items():
        if wildom_name in w_idx:
            col_map[sablon_idx + 1] = w_idx[wildom_name]  # 1-based excel col -> wildom idx

    # uj oszlopok hozzaadasa, ha meg nincsenek
    for header_text, wildom_name in NEW_COLUMNS.items():
        if wildom_name not in w_idx:
            continue
        if header_text in header_row:
            col = header_row.index(header_text) + 1
        else:
            max_col += 1
            col = max_col
            hc = ws.cell(row=5, column=col, value=header_text)
            hc.font = copy.copy(sample_header.font)
            hc.fill = copy.copy(sample_header.fill)
            hc.alignment = copy.copy(sample_header.alignment)
            hc.border = copy.copy(sample_header.border)
            ws.column_dimensions[hc.column_letter].width = ws.column_dimensions['D'].width
        col_map[col] = w_idx[wildom_name]

    unmapped = [h for h in w_headers[2:-1] if h and h not in COL_MAP_BY_NAME.values()
                and h not in NEW_COLUMNS.values() and h not in EXCLUDED_WILDOM_COLS]
    if unmapped:
        print("FIGYELEM - ismeretlen uj Wildom oszlop(ok), nincs sablon oszlop hozzajuk, "
              "ezeket NEM toltottuk ki:", unmapped)

    filled = 0
    for r in range(6, ws.max_row + 1):
        cikk = ws.cell(row=r, column=2).value
        if cikk is None:
            continue
        wildom_product = next((p for p, code in PRODUCT_MAP.items() if code == cikk), None)
        if wildom_product is None or wildom_product not in w_data:
            continue
        wrow = w_data[wildom_product]
        for col, widx in col_map.items():
            val = wrow[widx]
            cell = ws.cell(row=r, column=col)
            if val:
                cell.value = val
                filled += 1
            cell.font = copy.copy(sample_qty.font)
            cell.alignment = copy.copy(sample_qty.alignment)
            cell.border = copy.copy(sample_qty.border)

    out_path = sablon_path.rsplit(".", 1)[0] + f"_Cola_{new_sheet_name}.xlsx"
    wb_p.save(out_path)
    print(f"OK - '{new_sheet_name}' lap letrehozva, {filled} mennyiseg cella kitoltve.")
    print("Mentve:", out_path)


if __name__ == "__main__":
    main()
