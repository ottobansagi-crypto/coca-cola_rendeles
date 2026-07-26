"""
Wildom Cola-rendeles -> PizzaMe_megrendelo sablon automatikus feltoltese.

Hasznalat:
    python3 wildom_to_cola.py <wildom_export.xlsx> <pizzame_sablon.xlsx> [uj_lap_nev]

- <wildom_export.xlsx>: a Wildom "Beszallitoi megrendelesek" -> "Excel 2007+"
  exporttal letoltott fajl (pl. "megrendeles-Coca Cola-YYYYMMDD-HHMM.xlsx").
- <pizzame_sablon.xlsx>: a PizzaMe_megrendelo fajl (a legutobbi allapot,
  amibol uj lapot masolunk).
- [uj_lap_nev]: opcionalis, ha nem adod meg, a mai datum lesz (YYYYMMDD).

A script megkeresi a legfrissebb lapot a sablonban (a legnagyobb, csak
szambol allo lapnevet), lemasolja azt uj lapkent, es a Wildom exportbol
kitolti a mennyisegeket a rogzitett termek- es uzlet-parositas alapjan.

A termek- es uzletlista karbantartasa NEM ebben a fajlban tortenik: lasd a
../config/parositas.json fajlt. Ha a Coca-Cola cikkszamot/arat valtoztat,
vagy uzlet nyilik/zar/atnevezodik, azt a JSON-t kell szerkeszteni, nem ezt
a kodot.
"""
import sys
import re
import json
import copy
import datetime
from pathlib import Path
import openpyxl

CONFIG_PATH = Path(__file__).resolve().parent.parent / "config" / "parositas.json"


def load_config(path=CONFIG_PATH):
    with open(path, encoding="utf-8") as f:
        raw = json.load(f)

    product_map = {p["wildom_nev"]: p["cikkszam"] for p in raw["termek_parositas"]}
    col_map_by_name = {u["sablon_oszlop_index"]: u["wildom_nev"] for u in raw["uzlet_parositas"]}
    new_columns = {c["sablon_felirat"]: c["wildom_nev"] for c in raw["uj_oszlopok"]}
    excluded_wildom_cols = {e["wildom_nev"] for e in raw["kihagyott_wildom_oszlopok"]}

    return product_map, col_map_by_name, new_columns, excluded_wildom_cols


def find_latest_sheet(wb):
    candidates = [s for s in wb.sheetnames if re.fullmatch(r"\d{6,8}", s)]
    if not candidates:
        return wb.sheetnames[0]
    return sorted(candidates, key=lambda s: int(s))[-1]


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    wildom_path, sablon_path = sys.argv[1], sys.argv[2]
    new_sheet_name = sys.argv[3] if len(sys.argv) > 3 else datetime.date.today().strftime("%Y%m%d")

    product_map, col_map_by_name, new_columns, excluded_wildom_cols = load_config()

    wb_w = openpyxl.load_workbook(wildom_path, data_only=True)
    ws_w = wb_w[wb_w.sheetnames[0]]
    w_headers = [c.value for c in ws_w[1]]
    w_idx = {name: i for i, name in enumerate(w_headers)}
    w_data = {row[0]: row for row in ws_w.iter_rows(min_row=2, max_row=ws_w.max_row, values_only=True) if row[0]}

    wb_p = openpyxl.load_workbook(sablon_path)
    src_name = find_latest_sheet(wb_p)
    src = wb_p[src_name]

    if new_sheet_name in wb_p.sheetnames:
        del wb_p[new_sheet_name]
    ws = wb_p.copy_worksheet(src)
    ws.title = new_sheet_name

    max_col = ws.max_column
    sample_header = ws.cell(row=5, column=4)
    sample_qty = ws.cell(row=6, column=4)

    # meglevo oszlopok fejlistaja a masolt lapon
    header_row = [ws.cell(row=5, column=c).value for c in range(1, max_col + 1)]

    col_map = {}
    for sablon_idx, wildom_name in col_map_by_name.items():
        if wildom_name in w_idx:
            col_map[sablon_idx + 1] = w_idx[wildom_name]  # 1-based excel col -> wildom idx

    # uj oszlopok hozzaadasa, ha meg nincsenek
    for header_text, wildom_name in new_columns.items():
        if wildom_name not in w_idx:
            continue
        if header_text in header_row:
            col = header_row.index(header_text) + 1
        else:
            max_col += 1
            col = max_col
            hc = ws.cell(row=5, column=col, value=header_text)
            hc.font = copy.copy(sample_header.font)
            hc.fill = copy.copy(sample_header.fill)
            hc.alignment = copy.copy(sample_header.alignment)
            hc.border = copy.copy(sample_header.border)
            ws.column_dimensions[hc.column_letter].width = ws.column_dimensions['D'].width
        col_map[col] = w_idx[wildom_name]

    unmapped = [h for h in w_headers[2:-1] if h and h not in col_map_by_name.values()
                and h not in new_columns.values() and h not in excluded_wildom_cols]
    if unmapped:
        print("FIGYELEM - ismeretlen uj Wildom oszlop(ok), nincs sablon oszlop hozzajuk, "
              "ezeket NEM toltottuk ki:", unmapped)

    filled = 0
    for r in range(6, ws.max_row + 1):
        cikk = ws.cell(row=r, column=2).value
        if cikk is None:
            continue
        wildom_product = next((p for p, code in product_map.items() if code == cikk), None)
        if wildom_product is None or wildom_product not in w_data:
            continue
        wrow = w_data[wildom_product]
        for col, widx in col_map.items():
            val = wrow[widx]
            cell = ws.cell(row=r, column=col)
            if val:
                cell.value = val
                filled += 1
            cell.font = copy.copy(sample_qty.font)
            cell.alignment = copy.copy(sample_qty.alignment)
            cell.border = copy.copy(sample_qty.border)

    out_path = sablon_path.rsplit(".", 1)[0] + f"_Cola_{new_sheet_name}.xlsx"
    wb_p.save(out_path)
    print(f"OK - '{new_sheet_name}' lap letrehozva, {filled} mennyiseg cella kitoltve.")
    print("Mentve:", out_path)


if __name__ == "__main__":
    main()
